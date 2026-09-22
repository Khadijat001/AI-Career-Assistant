from io import BytesIO

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer
)
from reportlab.lib.units import inch


# ==========================================
# CREATE WORD DOCUMENT
# ==========================================

def create_docx(cv_text):

    document = Document()

    # Add title
    title = document.add_paragraph()
    title.alignment = TA_CENTER

    run = title.add_run("TAILORED CV")
    run.bold = True
    run.font.size = document.styles["Title"].font.size

    # Add CV content
    for line in cv_text.split("\n"):

        line = line.strip()

        if not line:
            document.add_paragraph()
            continue

        paragraph = document.add_paragraph()

        # Make section headings bold
        if line.isupper():
            run = paragraph.add_run(line)
            run.bold = True

        else:
            paragraph.add_run(line)

    # Save document to memory
    output = BytesIO()

    document.save(output)

    output.seek(0)

    return output.getvalue()


# ==========================================
# CREATE PDF DOCUMENT
# ==========================================

def create_pdf(cv_text):

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=0.6 * inch,
        leftMargin=0.6 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]
    title_style.alignment = TA_CENTER

    heading_style = styles["Heading2"]

    body_style = styles["BodyText"]

    story = []

    # PDF title
    story.append(
        Paragraph("TAILORED CV", title_style)
    )

    story.append(
        Spacer(1, 0.2 * inch)
    )

    # Add CV content
    for line in cv_text.split("\n"):

        line = line.strip()

        if not line:
            story.append(
                Spacer(1, 0.08 * inch)
            )
            continue

        # Section headings
        if line.isupper():

            story.append(
                Paragraph(
                    line,
                    heading_style
                )
            )

        else:

            story.append(
                Paragraph(
                    line,
                    body_style
                )
            )

    # Build PDF
    document.build(story)

    output.seek(0)

    return output.getvalue()